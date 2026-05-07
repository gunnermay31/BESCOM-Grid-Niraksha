from __future__ import annotations

import calendar
import re
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
LOAD_CURVE_DIR = ROOT / "data" / "raw" / "load_curve" / "2024"
OUTPUT_DIR = ROOT / "outputs"
OUTPUT_XLSX = OUTPUT_DIR / "kptcl_load_curve_2024_merged.xlsx"

FILE_RE = re.compile(r"^D(?P<day>\d{2})(?P<month>[A-Z]{3})(?P<year>\d{4})\.xls$", re.I)
MONTH_LOOKUP = {abbr.upper(): index for index, abbr in enumerate(calendar.month_abbr) if abbr}


def file_date(path: Path) -> datetime:
    match = FILE_RE.match(path.name)
    if not match:
        raise ValueError(f"Unexpected load-curve file name: {path.name}")
    return datetime(
        int(match.group("year")),
        MONTH_LOOKUP[match.group("month").upper()],
        int(match.group("day")),
    )


def to_number(value):
    if pd.isna(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def read_hourly_rows(path: Path) -> list[dict]:
    day = file_date(path)
    raw = pd.read_excel(path, sheet_name="LOAD CURVE", header=None, engine="xlrd")
    rows: list[dict] = []

    for _, source in raw.iterrows():
        hour_value = source.iloc[34] if len(source) > 34 else None
        load_value = source.iloc[35] if len(source) > 35 else None
        frequency_value = source.iloc[36] if len(source) > 36 else None
        previous_load_value = source.iloc[38] if len(source) > 38 else None

        hour = to_number(hour_value)
        load_mw = to_number(load_value)
        if hour is None or load_mw is None:
            continue
        if hour < 0 or hour > 23 or hour != int(hour):
            continue

        hour = int(hour)
        timestamp = day.replace(hour=hour)
        frequency_hz = to_number(frequency_value)
        previous_load_mw = to_number(previous_load_value)
        yoy_change_mw = load_mw - previous_load_mw if previous_load_mw not in (None, 0) else None
        yoy_change_pct = yoy_change_mw / previous_load_mw if yoy_change_mw is not None else None

        rows.append(
            {
                "date": day.date(),
                "timestamp": timestamp,
                "year": day.year,
                "month": day.month,
                "month_name": day.strftime("%B"),
                "day": day.day,
                "day_of_year": int(day.strftime("%j")),
                "day_of_week": day.strftime("%A"),
                "is_weekend": day.weekday() >= 5,
                "hour": hour,
                "time_block": f"{hour:02d}:00-{(hour + 1) % 24:02d}:00",
                "load_mw": load_mw,
                "frequency_hz": frequency_hz,
                "previous_year_load_mw": previous_load_mw,
                "yoy_load_change_mw": yoy_change_mw,
                "yoy_load_change_pct": yoy_change_pct,
                "source_file": path.name,
            }
        )

    return rows


def build_dataset() -> pd.DataFrame:
    all_rows: list[dict] = []
    for path in sorted(LOAD_CURVE_DIR.glob("*.xls"), key=file_date):
        all_rows.extend(read_hourly_rows(path))

    df = pd.DataFrame(all_rows).sort_values(["timestamp"]).reset_index(drop=True)
    if df.empty:
        raise RuntimeError(f"No hourly load-curve rows found in {LOAD_CURVE_DIR}")

    day_group = df.groupby("date")["load_mw"]
    df["daily_peak_load_mw"] = day_group.transform("max")
    df["daily_min_load_mw"] = day_group.transform("min")
    df["daily_avg_load_mw"] = day_group.transform("mean")
    df["is_daily_peak_hour"] = df["load_mw"].eq(df["daily_peak_load_mw"])
    df["is_daily_min_hour"] = df["load_mw"].eq(df["daily_min_load_mw"])
    df["load_factor_vs_daily_peak"] = df["load_mw"] / df["daily_peak_load_mw"]
    df["missing_date_note"] = ""

    expected_dates = {d.date() for d in pd.date_range("2024-01-01", "2024-12-31", freq="D")}
    actual_dates = set(df["date"])
    missing_dates = sorted(expected_dates - actual_dates)
    if missing_dates:
        note = "Missing source file for " + ", ".join(d.isoformat() for d in missing_dates)
        df.loc[df.index[0], "missing_date_note"] = note

    ordered = [
        "date",
        "timestamp",
        "year",
        "month",
        "month_name",
        "day",
        "day_of_year",
        "day_of_week",
        "is_weekend",
        "hour",
        "time_block",
        "load_mw",
        "frequency_hz",
        "previous_year_load_mw",
        "yoy_load_change_mw",
        "yoy_load_change_pct",
        "daily_peak_load_mw",
        "daily_min_load_mw",
        "daily_avg_load_mw",
        "is_daily_peak_hour",
        "is_daily_min_hour",
        "load_factor_vs_daily_peak",
        "source_file",
        "missing_date_note",
    ]
    return df[ordered]


def style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb.active
    ws.title = "Load Curve 2024"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 12,
        "B": 20,
        "E": 12,
        "H": 13,
        "K": 14,
        "L": 12,
        "M": 13,
        "N": 20,
        "O": 18,
        "P": 18,
        "Q": 18,
        "R": 17,
        "S": 17,
        "V": 20,
        "W": 17,
        "X": 44,
    }
    for idx, _ in enumerate(ws[1], start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(letter, 11)

    for row in ws.iter_rows(min_row=2, min_col=1, max_col=ws.max_column):
        row[0].number_format = "yyyy-mm-dd"
        row[1].number_format = "yyyy-mm-dd hh:mm"
        for idx in [11, 12, 13, 14, 16, 17, 18]:
            row[idx].number_format = "0.00"
        row[15].number_format = "0.00%"
        row[21].number_format = "0.00%"

    wb.save(path)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    df = build_dataset()
    df.to_excel(OUTPUT_XLSX, index=False, engine="openpyxl")
    style_workbook(OUTPUT_XLSX)
    print(f"Output: {OUTPUT_XLSX}")
    print(f"Rows: {len(df)}")
    print(f"Unique dates: {df['date'].nunique()}")
    print(f"Missing date notes: {df['missing_date_note'].dropna().astype(bool).sum()}")


if __name__ == "__main__":
    main()
