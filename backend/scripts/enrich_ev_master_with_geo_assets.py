from __future__ import annotations

import json
import math
import re
from datetime import datetime
from difflib import get_close_matches
from pathlib import Path

import pandas as pd
import requests
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUTAGE_PDF_DIR = ROOT / "data" / "raw" / "outages" / "2024"
LOAD_CURVE_DIR = ROOT / "data" / "raw" / "load_curve" / "2024"
EXTERNAL_DIR = ROOT / "data" / "external"
OUTPUT_DIR = ROOT / "outputs"

MASTER_INPUT_CSV = OUTPUT_DIR / "ev_grid_eda_master_2024.csv"
MASTER_INPUT_XLSX = OUTPUT_DIR / "ev_grid_eda_master_2024.xlsx"

OUT_EVENT_OUTPUT_CSV = OUTPUT_DIR / "ev_grid_outage_events_enriched_2024.csv"
OUT_EVENT_OUTPUT_XLSX = OUTPUT_DIR / "ev_grid_outage_events_enriched_2024.xlsx"
MASTER_OUTPUT_CSV = OUTPUT_DIR / "ev_grid_eda_master_2024_enriched.csv"
MASTER_OUTPUT_XLSX = OUTPUT_DIR / "ev_grid_eda_master_2024_enriched.xlsx"

SUBSTATION_CACHE = EXTERNAL_DIR / "osm_karnataka_substations.json"
EV_CACHE = EXTERNAL_DIR / "osm_karnataka_ev_charging.json"

OVERPASS_URL = "https://overpass-api.de/api/interpreter"
KARNATAKA_BBOX = "(11.4,73.8,18.7,78.7)"

ROW_START_RE = re.compile(r"(?<!\d)(?P<serial>\d{1,6})\s+(?P<date>\d{2}-[A-Za-z]{3}-\d{2})\b")
ZONE_RE = re.compile(r"^(?P<zone>[A-Z][A-Z /&().-]{2,}?)\s+(?P<voltage>\d{2,3})\s+(?P<rest>.*)$")
TIME_RE = re.compile(r"\b\d{1,2}:\d{2}(?::\d{2})?\b")
STATION_CODE_RE = re.compile(r"\b([A-Z][A-Z0-9]{1,20}(?: [A-Z0-9]{1,20}){0,4}_[0-9]{2,3})\b")


def norm_text(text: str) -> str:
    cleaned = re.sub(r"[^A-Z0-9 ]+", " ", text.upper())
    return re.sub(r"\s+", " ", cleaned).strip()


def safe_float(value):
    if pd.isna(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_hhmm(text: str | None) -> str | None:
    if not text:
        return None
    match = re.search(r"(\d{1,2}):(\d{2})", str(text))
    if not match:
        return None
    hour = int(match.group(1))
    minute = int(match.group(2))
    if hour == 24 and minute == 0:
        return "00:00"
    if 0 <= hour <= 23 and 0 <= minute <= 59:
        return f"{hour:02d}:{minute:02d}"
    return None


def haversine_km(lat1, lon1, lat2, lon2):
    r = 6371.0
    p1 = math.radians(lat1)
    p2 = math.radians(lat2)
    dlat = p2 - p1
    dlon = math.radians(lon2 - lon1)
    a = math.sin(dlat / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dlon / 2) ** 2
    return 2 * r * math.atan2(math.sqrt(a), math.sqrt(1 - a))


def query_overpass(query: str) -> dict:
    headers = {"User-Agent": "codex-ev-grid-enrichment/1.0"}
    response = requests.post(OVERPASS_URL, data=query.encode("utf-8"), headers=headers, timeout=240)
    response.raise_for_status()
    return response.json()


def fetch_or_load_osm_cache(cache_path: Path, query: str, kind: str) -> list[dict]:
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    if cache_path.exists():
        with cache_path.open("r", encoding="utf-8") as f:
            return json.load(f)

    payload = query_overpass(query)
    records: list[dict] = []
    for el in payload.get("elements", []):
        tags = el.get("tags", {}) or {}
        lat = el.get("lat")
        lon = el.get("lon")
        if lat is None or lon is None:
            center = el.get("center", {}) or {}
            lat = center.get("lat")
            lon = center.get("lon")
        if lat is None or lon is None:
            continue

        name = tags.get("name") or tags.get("name:en") or tags.get("ref") or f"{kind}_{el.get('type')}_{el.get('id')}"
        records.append(
            {
                "id": el.get("id"),
                "osm_type": el.get("type"),
                "name": str(name),
                "name_norm": norm_text(str(name)),
                "lat": float(lat),
                "lon": float(lon),
                "tags": tags,
            }
        )

    with cache_path.open("w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=True)
    return records


def load_substations() -> list[dict]:
    query = (
        "[out:json][timeout:180];"
        f"(node[\"power\"=\"substation\"]{KARNATAKA_BBOX};"
        f"way[\"power\"=\"substation\"]{KARNATAKA_BBOX};"
        f"relation[\"power\"=\"substation\"]{KARNATAKA_BBOX};);"
        "out center tags;"
    )
    return fetch_or_load_osm_cache(SUBSTATION_CACHE, query, "substation")


def load_ev_points() -> list[dict]:
    query = (
        "[out:json][timeout:180];"
        f"(node[\"amenity\"=\"charging_station\"]{KARNATAKA_BBOX};"
        f"way[\"amenity\"=\"charging_station\"]{KARNATAKA_BBOX};"
        f"relation[\"amenity\"=\"charging_station\"]{KARNATAKA_BBOX};);"
        "out center tags;"
    )
    return fetch_or_load_osm_cache(EV_CACHE, query, "ev")


def choose_best_substation(
    candidates: list[str],
    exact_index: dict[str, list[dict]],
    all_names: list[str],
    prefix_index: dict[str, list[dict]],
) -> dict | None:
    for raw in candidates:
        key = norm_text(raw)
        if not key:
            continue
        if key in exact_index:
            return exact_index[key][0]

        first_token = key.split(" ")[0]
        scoped = prefix_index.get(first_token, [])
        contains = [s for s in scoped if key in s["name_norm"] or s["name_norm"] in key]
        if contains:
            return sorted(contains, key=lambda s: abs(len(s["name_norm"]) - len(key)))[0]

        close = get_close_matches(key, all_names, n=1, cutoff=0.88)
        if close:
            return exact_index[close[0]][0]
    return None


def extract_candidate_locations(asset_text: str) -> list[str]:
    candidates: list[str] = []

    for token in STATION_CODE_RE.findall(asset_text.upper()):
        station = re.sub(r"_[0-9]{2,3}$", "", token).strip()
        station = re.sub(r"\s+", " ", station)
        if len(station) >= 3:
            candidates.append(station)

    if "-" in asset_text:
        parts = re.split(r"\s*-\s*", asset_text.upper())
        for part in parts[:3]:
            part = re.sub(r"\b(LINE|TRANSFORMER|STATION|NOT|AFFECTED|NO|KV|HV\d+)\b", " ", part)
            part = re.sub(r"\b\d+\b", " ", part)
            part = re.sub(r"\s+", " ", part).strip()
            if 3 <= len(part) <= 40:
                candidates.append(part)

    seen = set()
    deduped: list[str] = []
    for item in candidates:
        key = norm_text(item)
        if key and key not in seen:
            deduped.append(item)
            seen.add(key)
    return deduped


def classify_asset_type(text: str) -> str:
    upper = text.upper()
    if "TRANSFORMER" in upper or re.search(r"\bHV[0-9]\b", upper):
        return "transformer"
    if "LINE" in upper or "-" in upper:
        return "line"
    return "other"


def parse_outage_events() -> pd.DataFrame:
    rows: list[dict] = []
    for pdf in sorted(OUTAGE_PDF_DIR.glob("*.pdf")):
        reader = PdfReader(str(pdf))
        for page_num, page in enumerate(reader.pages, start=1):
            text = re.sub(r"\s+", " ", (page.extract_text() or "").replace("\xa0", " ")).strip()
            matches = list(ROW_START_RE.finditer(text))
            for idx, match in enumerate(matches):
                start = match.start()
                end = matches[idx + 1].start() if idx + 1 < len(matches) else len(text)
                chunk = text[start:end].strip()
                head = ROW_START_RE.match(chunk)
                if not head:
                    continue

                date_raw = head.group("date")
                serial = int(head.group("serial"))
                body = chunk[head.end() :].strip()
                zone = ""
                voltage_class_kv = None
                asset_text = body

                zone_match = ZONE_RE.match(body)
                if zone_match:
                    zone = zone_match.group("zone").strip()
                    voltage_class_kv = int(zone_match.group("voltage"))
                    asset_text = zone_match.group("rest").strip()

                times = TIME_RE.findall(asset_text)
                opened_at = parse_hhmm(times[0]) if len(times) >= 1 else None
                closed_at = parse_hhmm(times[1]) if len(times) >= 2 else None

                duration_minutes = None
                duration_match = re.search(r"\b(\d{1,2}:\d{2}(?::\d{2})?)\b", asset_text)
                if duration_match:
                    hhmm = parse_hhmm(duration_match.group(1))
                    if hhmm:
                        hh, mm = map(int, hhmm.split(":"))
                        duration_minutes = hh * 60 + mm

                try:
                    date_obj = datetime.strptime(date_raw, "%d-%b-%y").date()
                except ValueError:
                    continue

                opened_hour = None
                event_hour_ts = None
                if opened_at:
                    opened_hour = int(opened_at.split(":")[0])
                    event_hour_ts = datetime.combine(date_obj, datetime.strptime(opened_at, "%H:%M").time()).replace(minute=0, second=0)

                asset_type = classify_asset_type(asset_text)
                mva_values = [float(x) for x in re.findall(r"(\d+(?:\.\d+)?)\s*MVA", asset_text.upper())]
                kv_values = [float(x) for x in re.findall(r"(\d{2,3})\s*KV", asset_text.upper())]
                transformer_capacity_mva = max(mva_values) if mva_values else None
                line_capacity_kv = max(kv_values) if kv_values else (float(voltage_class_kv) if asset_type == "line" and voltage_class_kv else None)
                candidates = extract_candidate_locations(asset_text)

                rows.append(
                    {
                        "source_pdf": pdf.name,
                        "source_page": page_num,
                        "serial_no": serial,
                        "date": date_obj,
                        "date_raw": date_raw,
                        "event_hour_ts": event_hour_ts,
                        "opened_hour": opened_hour,
                        "transmission_zone": zone,
                        "voltage_class_kv": voltage_class_kv,
                        "asset_station_text": asset_text,
                        "asset_type": asset_type,
                        "transformer_capacity_mva": transformer_capacity_mva if asset_type == "transformer" else None,
                        "line_capacity_kv": line_capacity_kv if asset_type == "line" else None,
                        "candidate_locations": "|".join(candidates),
                        "opened_at": opened_at,
                        "closed_at": closed_at,
                        "duration_minutes": duration_minutes,
                    }
                )
    df = pd.DataFrame(rows)
    return df


def parse_daily_resource_mix() -> pd.DataFrame:
    records: list[dict] = []

    thermal_keys = ("RTPS", "BTPS", "YTPS", "KUDGI", "NTPC", "FSTPP")
    hydro_keys = (
        "SHARAVATHY",
        "N. P. H",
        "VARAHI",
        "GERUSOPPA",
        "ALMATTI",
        "KADRA",
        "KODASALLY",
        "SUPA",
        "L.D.P.H",
        "BHADRA",
        "GHATAPRABA",
        "M.D.P.H",
        "M.G.H.E",
        "SIVASAMUDRA",
        "MUNIRABAD",
        "JURALA",
        "TB DAM SHARE",
    )
    solar_keys = ("SOLAR",)
    cgs_keys = ("NET CGS IMPORT", "RAILWAYS")

    for path in sorted(LOAD_CURVE_DIR.glob("*.xls")):
        file_match = re.match(r"^D(\d{2})([A-Z]{3})(\d{4})\.xls$", path.name, flags=re.I)
        if not file_match:
            continue
        day = int(file_match.group(1))
        mon = datetime.strptime(file_match.group(2).title(), "%b").month
        year = int(file_match.group(3))
        date_obj = datetime(year, mon, day).date()

        raw = pd.read_excel(path, sheet_name="LOAD CURVE", header=None, engine="xlrd")
        totals = {"thermal": 0.0, "hydro": 0.0, "solar": 0.0, "cgs_import": 0.0, "other": 0.0}

        for idx in range(min(len(raw), 70)):
            name = str(raw.iloc[idx, 0]) if idx < len(raw) else ""
            gen_mu = safe_float(raw.iloc[idx, 8] if raw.shape[1] > 8 else None)
            station_no = safe_float(raw.iloc[idx, 3] if raw.shape[1] > 3 else None)
            station_mw = safe_float(raw.iloc[idx, 4] if raw.shape[1] > 4 else None)
            if gen_mu is None or gen_mu < 0:
                continue
            if station_no is None or station_mw is None:
                continue
            name_u = name.upper()
            name_norm = norm_text(name_u)
            name_compact = name_norm.replace(" ", "")
            if any(skip in name_u for skip in ("STATIONS", "TOTAL", "SCHEDULE", "PROGRESSIVE", "AVAILABILITY", "CENTRAL GENERATOR OUTAGES")):
                continue

            if any(k in name_norm for k in solar_keys):
                totals["solar"] += gen_mu
            elif any(k in name_compact for k in thermal_keys):
                totals["thermal"] += gen_mu
            elif any(norm_text(k).replace(" ", "") in name_compact for k in hydro_keys):
                totals["hydro"] += gen_mu
            elif any(norm_text(k).replace(" ", "") in name_compact for k in cgs_keys):
                totals["cgs_import"] += gen_mu
            else:
                totals["other"] += gen_mu

        total_mu = sum(totals.values())
        dominant = max(totals, key=totals.get) if total_mu > 0 else "unknown"
        records.append(
            {
                "date": date_obj,
                "resource_total_mu_day": total_mu,
                "resource_thermal_mu_day": totals["thermal"],
                "resource_hydro_mu_day": totals["hydro"],
                "resource_solar_mu_day": totals["solar"],
                "resource_cgs_import_mu_day": totals["cgs_import"],
                "resource_other_mu_day": totals["other"],
                "resource_thermal_share_day": (totals["thermal"] / total_mu) if total_mu else None,
                "resource_hydro_share_day": (totals["hydro"] / total_mu) if total_mu else None,
                "resource_solar_share_day": (totals["solar"] / total_mu) if total_mu else None,
                "resource_cgs_import_share_day": (totals["cgs_import"] / total_mu) if total_mu else None,
                "dominant_power_resource_day": dominant,
            }
        )
    return pd.DataFrame(records)


def style_sheet(path: Path, title: str):
    wb = load_workbook(path)
    ws = wb.active
    ws.title = title
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for idx, _ in enumerate(ws[1], start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = 18
    wb.save(path)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    EXTERNAL_DIR.mkdir(parents=True, exist_ok=True)

    if MASTER_INPUT_CSV.exists():
        master = pd.read_csv(MASTER_INPUT_CSV)
        master["date"] = pd.to_datetime(master["date"], format="%Y-%m-%d", errors="coerce").dt.date
        master["timestamp"] = pd.to_datetime(master["timestamp"], format="%Y-%m-%d %H:%M:%S", errors="coerce")
    elif MASTER_INPUT_XLSX.exists():
        master = pd.read_excel(MASTER_INPUT_XLSX)
        master["date"] = pd.to_datetime(master["date"]).dt.date
        master["timestamp"] = pd.to_datetime(master["timestamp"])
    else:
        raise FileNotFoundError("Run build_ev_eda_master_2024.py first to create the master dataset.")

    substations = load_substations()
    ev_points = load_ev_points()
    exact_index: dict[str, list[dict]] = {}
    prefix_index: dict[str, list[dict]] = {}
    for s in substations:
        exact_index.setdefault(s["name_norm"], []).append(s)
        token = s["name_norm"].split(" ")[0] if s["name_norm"] else ""
        if token:
            prefix_index.setdefault(token, []).append(s)
    all_names = list(exact_index.keys())

    outage = parse_outage_events()
    match_cache: dict[str, dict | None] = {}

    # Match outage candidates to substations and derive event coordinates.
    geo_rows = []
    for _, row in outage.iterrows():
        candidates = [c for c in str(row["candidate_locations"]).split("|") if c and c.lower() != "nan"]
        cache_key = "||".join(norm_text(c) for c in candidates if c)
        if cache_key in match_cache:
            best = match_cache[cache_key]
        else:
            best = choose_best_substation(candidates, exact_index, all_names, prefix_index)
            match_cache[cache_key] = best

        lat = best["lat"] if best else None
        lon = best["lon"] if best else None
        loc_name = best["name"] if best else None
        match_method = "substation_name_match" if best else "no_match"

        nearest_name = None
        nearest_lat = None
        nearest_lon = None
        nearest_km = None
        if lat is not None and lon is not None and ev_points:
            nearest = min(ev_points, key=lambda ev: haversine_km(lat, lon, ev["lat"], ev["lon"]))
            nearest_name = nearest["name"]
            nearest_lat = nearest["lat"]
            nearest_lon = nearest["lon"]
            nearest_km = haversine_km(lat, lon, nearest_lat, nearest_lon)

        geo_rows.append(
            {
                "event_geo_lat": lat,
                "event_geo_lon": lon,
                "event_location_name_matched": loc_name,
                "event_location_match_method": match_method,
                "nearest_ev_point_name": nearest_name,
                "nearest_ev_point_lat": nearest_lat,
                "nearest_ev_point_lon": nearest_lon,
                "nearest_ev_point_distance_km": nearest_km,
            }
        )

    outage = pd.concat([outage.reset_index(drop=True), pd.DataFrame(geo_rows)], axis=1)

    resource_mix = parse_daily_resource_mix()
    outage = outage.merge(resource_mix, on="date", how="left")
    outage["likely_power_source_for_transformer"] = outage["dominant_power_resource_day"].where(outage["asset_type"] == "transformer")

    outage.to_csv(OUT_EVENT_OUTPUT_CSV, index=False)
    outage.to_excel(OUT_EVENT_OUTPUT_XLSX, index=False)
    style_sheet(OUT_EVENT_OUTPUT_XLSX, "Outage_Events_Enriched")

    # Aggregate event-level additions into hourly features and merge into master.
    outage_hourly = outage.dropna(subset=["event_hour_ts"]).copy()
    hourly = (
        outage_hourly.groupby("event_hour_ts")
        .agg(
            transformer_outages_started_hour=("asset_type", lambda s: (s == "transformer").sum()),
            line_outages_started_hour=("asset_type", lambda s: (s == "line").sum()),
            avg_transformer_capacity_mva_hour=("transformer_capacity_mva", "mean"),
            max_transformer_capacity_mva_hour=("transformer_capacity_mva", "max"),
            avg_line_capacity_kv_hour=("line_capacity_kv", "mean"),
            max_line_capacity_kv_hour=("line_capacity_kv", "max"),
            outage_location_centroid_lat_hour=("event_geo_lat", "mean"),
            outage_location_centroid_lon_hour=("event_geo_lon", "mean"),
            nearest_ev_distance_km_hour=("nearest_ev_point_distance_km", "mean"),
            nearest_ev_distance_km_min_hour=("nearest_ev_point_distance_km", "min"),
            nearest_ev_point_name_hour=("nearest_ev_point_name", lambda s: s.dropna().iloc[0] if len(s.dropna()) else None),
            likely_resource_for_transformer_hour=("likely_power_source_for_transformer", lambda s: s.dropna().mode().iloc[0] if len(s.dropna()) else None),
        )
        .reset_index()
    )

    master["timestamp"] = pd.to_datetime(master["timestamp"])
    master["date"] = pd.to_datetime(master["date"]).dt.date
    hourly["event_hour_ts"] = pd.to_datetime(hourly["event_hour_ts"])
    master = master.merge(hourly, left_on="timestamp", right_on="event_hour_ts", how="left").drop(columns=["event_hour_ts"], errors="ignore")
    master = master.merge(resource_mix, on="date", how="left", suffixes=("", "_daily_resource"))

    master.to_csv(MASTER_OUTPUT_CSV, index=False)
    master.to_excel(MASTER_OUTPUT_XLSX, index=False)
    style_sheet(MASTER_OUTPUT_XLSX, "EV_Grid_EDA_Enriched")

    print(f"Outage events enriched rows: {len(outage)}")
    print(f"Master enriched rows: {len(master)}")
    print(f"Substations loaded: {len(substations)}")
    print(f"EV points loaded: {len(ev_points)}")
    print(f"Event coordinates coverage: {outage['event_geo_lat'].notna().mean():.2%}")
    print(f"Nearest EV coverage: {outage['nearest_ev_point_distance_km'].notna().mean():.2%}")
    print(f"Output outage csv: {OUT_EVENT_OUTPUT_CSV}")
    print(f"Output master csv: {MASTER_OUTPUT_CSV}")


if __name__ == "__main__":
    main()
