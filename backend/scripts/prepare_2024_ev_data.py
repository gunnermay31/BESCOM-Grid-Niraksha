from __future__ import annotations

import calendar
import re
import urllib.error
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
OUTAGE_DIR = ROOT / "data" / "raw" / "outages" / "2024"
LOAD_CURVE_DIR = ROOT / "data" / "raw" / "load_curve" / "2024"
OUTPUT_DIR = ROOT / "outputs"
OUTPUT_XLSX = OUTPUT_DIR / "karnataka_ev_2024_outages.xlsx"
LOAD_URL = "https://loadcurve.kptcl.net/LoadCurveUpload/data/{name}.xls"

ROW_START_RE = re.compile(r"^(\d+)\s+(\d{2}-[A-Za-z]{3}-\d{2})\s+(.*)$")
ZONE_RE = re.compile(r"^(?P<zone>[A-Z][A-Z /&.-]*?)\s*(?P<voltage>\d{2,3})\s*(?P<rest>.*)$")
TIME_RE = re.compile(r"\b\d{1,2}:\d{2}\b")


@dataclass
class ParsedRow:
    source_pdf: str
    page: int
    serial_no: int
    outage_date: str
    transmission_zone: str
    voltage_class: str
    detail_text: str
    opened_at: str
    closed_at: str
    duration: str
    remarks: str


def iter_dates(year: int):
    day = date(year, 1, 1)
    while day.year == year:
        yield day
        day += timedelta(days=1)


def download_load_curves() -> list[dict[str, str]]:
    LOAD_CURVE_DIR.mkdir(parents=True, exist_ok=True)

    def fetch(day: date) -> dict[str, str]:
        file_stem = f"D{day.day:02d}{calendar.month_abbr[day.month].upper()}{day.year}"
        target = LOAD_CURVE_DIR / f"{file_stem}.xls"
        url = LOAD_URL.format(name=file_stem)

        if target.exists() and target.stat().st_size > 0:
            return {"date": day.isoformat(), "file": target.name, "status": "exists", "bytes": str(target.stat().st_size), "url": url}

        try:
            request = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
            with urllib.request.urlopen(request, timeout=8) as response:
                content = response.read()
            if not (content.startswith(b"\xd0\xcf\x11\xe0") or content.startswith(b"PK")):
                return {"date": day.isoformat(), "file": target.name, "status": "non_excel_response", "bytes": str(len(content)), "url": url}
            target.write_bytes(content)
            return {"date": day.isoformat(), "file": target.name, "status": "downloaded", "bytes": str(len(content)), "url": url}
        except urllib.error.HTTPError as exc:
            return {"date": day.isoformat(), "file": target.name, "status": f"http_{exc.code}", "bytes": "0", "url": url}
        except Exception as exc:
            return {"date": day.isoformat(), "file": target.name, "status": type(exc).__name__, "bytes": "0", "url": url}

    statuses: list[dict[str, str]] = []
    with ThreadPoolExecutor(max_workers=12) as executor:
        futures = [executor.submit(fetch, day) for day in iter_dates(2024)]
        for future in as_completed(futures):
            statuses.append(future.result())

    return sorted(statuses, key=lambda item: item["date"])


def clean_line(line: str) -> str:
    return re.sub(r"\s+", " ", line).strip()


def parse_outage_text(source_pdf: str, page_num: int, text: str) -> tuple[list[ParsedRow], list[list[str]]]:
    rows: list[ParsedRow] = []
    raw_lines: list[list[str]] = []

    current: str | None = None
    for original in text.splitlines():
        line = clean_line(original)
        if not line:
            continue
        raw_lines.append([source_pdf, page_num, line])
        if line.startswith("SL NO.") or line in {"CLASSLINE/TRANSFORMER AFFECTED STATIONSOPENED_AT", "(A)CLOSED_AT", "(B)DURATION REMARKS"}:
            continue
        if ROW_START_RE.match(line):
            if current:
                rows.append(parse_row(source_pdf, page_num, current))
            current = line
        elif current:
            current = f"{current} {line}"

    if current:
        rows.append(parse_row(source_pdf, page_num, current))
    return rows, raw_lines


def parse_row(source_pdf: str, page_num: int, line: str) -> ParsedRow:
    match = ROW_START_RE.match(line)
    if not match:
        return ParsedRow(source_pdf, page_num, 0, "", "", "", line, "", "", "", "")

    serial_no = int(match.group(1))
    outage_date = match.group(2)
    remainder = match.group(3)
    zone = voltage = ""
    detail = remainder

    zone_match = ZONE_RE.match(remainder)
    if zone_match:
        zone = clean_line(zone_match.group("zone"))
        voltage = zone_match.group("voltage")
        detail = zone_match.group("rest").strip()

    times = list(TIME_RE.finditer(detail))
    opened = closed = duration = remarks = ""
    if len(times) >= 2:
        opened = times[-2].group(0)
        closed = times[-1].group(0)
        before = detail[: times[-2].start()].strip()
        after = detail[times[-1].end() :].strip()
        dur_match = re.match(r"^(\d{1,2}:\d{2})\s*(.*)$", after)
        if dur_match:
            duration = dur_match.group(1)
            remarks = dur_match.group(2).strip()
        else:
            remarks = after
        detail = before
    elif len(times) == 1:
        opened = times[0].group(0)
        detail = detail[: times[0].start()].strip() + " " + detail[times[0].end() :].strip()

    return ParsedRow(source_pdf, page_num, serial_no, outage_date, zone, voltage, clean_line(detail), opened, closed, duration, clean_line(remarks))


def extract_outages() -> tuple[list[ParsedRow], list[list[str]], list[dict[str, str]]]:
    all_rows: list[ParsedRow] = []
    all_raw: list[list[str]] = []
    coverage: list[dict[str, str]] = []

    for pdf in sorted(OUTAGE_DIR.glob("*.pdf")):
        reader = PdfReader(str(pdf))
        start_count = len(all_rows)
        for index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            rows, raw_lines = parse_outage_text(pdf.name, index, text)
            all_rows.extend(rows)
            all_raw.extend(raw_lines)
        coverage.append({"source_pdf": pdf.name, "pages": str(len(reader.pages)), "parsed_rows": str(len(all_rows) - start_count)})

    return all_rows, all_raw, coverage


def style_sheet(ws):
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value or "")) for cell in column_cells[:200])
        ws.column_dimensions[letter].width = min(max(max_len + 2, 10), 55)


def write_workbook(outage_rows: list[ParsedRow], raw_lines: list[list[str]], coverage: list[dict[str, str]], load_statuses: list[dict[str, str]]):
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    ws = wb.active
    ws.title = "Outage Parsed"
    headers = [
        "source_pdf",
        "page",
        "serial_no",
        "date",
        "transmission_zone",
        "voltage_class",
        "line_transformer_affected",
        "opened_at",
        "closed_at",
        "duration",
        "remarks",
    ]
    ws.append(headers)
    for row in outage_rows:
        ws.append([
            row.source_pdf,
            row.page,
            row.serial_no,
            row.outage_date,
            row.transmission_zone,
            row.voltage_class,
            row.detail_text,
            row.opened_at,
            row.closed_at,
            row.duration,
            row.remarks,
        ])
    style_sheet(ws)

    raw_ws = wb.create_sheet("Outage Raw Lines")
    raw_ws.append(["source_pdf", "page", "raw_line"])
    for raw in raw_lines:
        raw_ws.append(raw)
    style_sheet(raw_ws)

    cov_ws = wb.create_sheet("Source Coverage")
    cov_ws.append(["source_pdf", "pages", "parsed_rows"])
    for item in coverage:
        cov_ws.append([item["source_pdf"], item["pages"], item["parsed_rows"]])
    style_sheet(cov_ws)

    load_ws = wb.create_sheet("Load Curve Downloads")
    load_ws.append(["date", "file", "status", "bytes", "url"])
    for item in load_statuses:
        load_ws.append([item["date"], item["file"], item["status"], item["bytes"], item["url"]])
    style_sheet(load_ws)

    wb.save(OUTPUT_XLSX)


def main():
    load_statuses = download_load_curves()
    outage_rows, raw_lines, coverage = extract_outages()
    write_workbook(outage_rows, raw_lines, coverage, load_statuses)

    downloaded = sum(1 for item in load_statuses if item["status"] in {"downloaded", "exists"})
    failed = len(load_statuses) - downloaded
    print(f"Workbook: {OUTPUT_XLSX}")
    print(f"Outage PDFs: {len(coverage)}")
    print(f"Parsed outage rows: {len(outage_rows)}")
    print(f"Raw outage lines: {len(raw_lines)}")
    print(f"Load curve files ready: {downloaded}; failed: {failed}")


if __name__ == "__main__":
    main()
