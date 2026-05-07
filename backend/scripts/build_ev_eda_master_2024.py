from __future__ import annotations

import calendar
import re
from datetime import datetime, timedelta
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
LOAD_CURVE_DIR = ROOT / "data" / "raw" / "load_curve" / "2024"
OUTAGE_PDF_DIR = ROOT / "data" / "raw" / "outages" / "2024"
OUTPUT_DIR = ROOT / "outputs"
OUTPUT_XLSX = OUTPUT_DIR / "ev_grid_eda_master_2024.xlsx"
OUTPUT_CSV = OUTPUT_DIR / "ev_grid_eda_master_2024.csv"

LOAD_FILE_RE = re.compile(r"^D(?P<day>\d{2})(?P<month>[A-Z]{3})(?P<year>\d{4})\.xls$", re.I)
ROW_START_RE = re.compile(r"(?<!\d)(?P<serial>\d{1,6})\s+(?P<date>\d{2}-[A-Za-z]{3}-\d{2})\b")
TIME_RE = re.compile(r"\b\d{1,2}:\d{2}(?::\d{2})?\b")
ZONE_RE = re.compile(r"^(?P<zone>[A-Z][A-Z /&().-]{2,}?)\s+(?P<voltage>\d{2,3})\s+(?P<rest>.*)$")

MONTH_LOOKUP = {abbr.upper(): idx for idx, abbr in enumerate(calendar.month_abbr) if abbr}


def parse_load_file_date(path: Path) -> datetime:
    match = LOAD_FILE_RE.match(path.name)
    if not match:
        raise ValueError(f"Unexpected load-curve file name: {path.name}")
    return datetime(
        int(match.group("year")),
        MONTH_LOOKUP[match.group("month").upper()],
        int(match.group("day")),
    )


def to_float(value) -> float | None:
    if pd.isna(value):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def normalize_time(value: str | None) -> str | None:
    if not value:
        return None
    match = re.search(r"(\d{1,2}:\d{2})(?::\d{2})?", str(value))
    return match.group(1) if match else None


def parse_time_to_datetime(base_date, hhmm: str | None) -> datetime | None:
    if hhmm is None:
        return None
    match = re.match(r"^(\d{1,2}):(\d{2})$", hhmm)
    if not match:
        return None
    hour = int(match.group(1))
    minute = int(match.group(2))
    day_offset = 0
    if hour == 24 and minute == 0:
        hour = 0
        day_offset = 1
    if hour < 0 or hour > 23 or minute < 0 or minute > 59:
        return None
    return datetime.combine(base_date, datetime.strptime(f"{hour:02d}:{minute:02d}", "%H:%M").time()) + timedelta(days=day_offset)


def parse_duration_minutes(value: str | None) -> float | None:
    if not value:
        return None
    match = re.search(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", value)
    if not match:
        return None
    hours = int(match.group(1))
    minutes = int(match.group(2))
    seconds = int(match.group(3) or 0)
    return hours * 60 + minutes + seconds / 60.0


def build_load_hourly_dataset() -> pd.DataFrame:
    rows: list[dict] = []

    for file_path in sorted(LOAD_CURVE_DIR.glob("*.xls"), key=parse_load_file_date):
        day = parse_load_file_date(file_path)
        raw = pd.read_excel(file_path, sheet_name="LOAD CURVE", header=None, engine="xlrd")

        for _, source in raw.iterrows():
            hour = to_float(source.iloc[34] if len(source) > 34 else None)
            load_2024 = to_float(source.iloc[35] if len(source) > 35 else None)
            frequency = to_float(source.iloc[36] if len(source) > 36 else None)
            load_2023 = to_float(source.iloc[38] if len(source) > 38 else None)

            if hour is None or load_2024 is None:
                continue
            if hour < 0 or hour > 23 or hour != int(hour):
                continue

            hour_int = int(hour)
            timestamp = day.replace(hour=hour_int)
            delta_mw = load_2024 - load_2023 if load_2023 not in (None, 0) else None
            delta_pct = (delta_mw / load_2023) if delta_mw is not None else None

            rows.append(
                {
                    "date": day.date(),
                    "timestamp": timestamp,
                    "year": day.year,
                    "month": day.month,
                    "month_name": day.strftime("%B"),
                    "quarter": f"Q{((day.month - 1) // 3) + 1}",
                    "day": day.day,
                    "day_of_year": int(day.strftime("%j")),
                    "day_of_week": day.strftime("%A"),
                    "is_weekend": day.weekday() >= 5,
                    "hour": hour_int,
                    "time_block": f"{hour_int:02d}:00-{(hour_int + 1) % 24:02d}:00",
                    "load_2024_mw": load_2024,
                    "load_2023_mw": load_2023,
                    "frequency_hz": frequency,
                    "load_delta_2024_vs_2023_mw": delta_mw,
                    "load_delta_2024_vs_2023_pct": delta_pct,
                    "load_curve_source_file": file_path.name,
                }
            )

    df = pd.DataFrame(rows).sort_values("timestamp").reset_index(drop=True)
    if df.empty:
        raise RuntimeError(f"No load-curve rows parsed from: {LOAD_CURVE_DIR}")

    day_group = df.groupby("date")["load_2024_mw"]
    df["daily_peak_load_2024_mw"] = day_group.transform("max")
    df["daily_min_load_2024_mw"] = day_group.transform("min")
    df["daily_avg_load_2024_mw"] = day_group.transform("mean")
    df["load_factor_vs_daily_peak"] = df["load_2024_mw"] / df["daily_peak_load_2024_mw"]
    df["is_daily_peak_hour"] = df["load_2024_mw"] == df["daily_peak_load_2024_mw"]
    df["is_daily_min_hour"] = df["load_2024_mw"] == df["daily_min_load_2024_mw"]
    df["ramp_1h_mw"] = df["load_2024_mw"].diff()
    df.loc[df["date"] != df["date"].shift(1), "ramp_1h_mw"] = np.nan
    df["load_percentile_within_day"] = df.groupby("date")["load_2024_mw"].rank(pct=True, method="average")

    expected_dates = pd.date_range("2024-01-01", "2024-12-31", freq="D").date
    available_dates = set(df["date"])
    missing_dates = [d.isoformat() for d in expected_dates if d not in available_dates]
    df["load_data_coverage_note"] = ""
    if missing_dates:
        df.loc[df.index[0], "load_data_coverage_note"] = "Missing load file for: " + ", ".join(missing_dates)

    return df


def parse_outage_chunks(text: str) -> list[tuple[int, str, str]]:
    clean = re.sub(r"\s+", " ", text.replace("\xa0", " ")).strip()
    matches = list(ROW_START_RE.finditer(clean))
    chunks: list[tuple[int, str, str]] = []
    for idx, match in enumerate(matches):
        start = match.start()
        end = matches[idx + 1].start() if idx + 1 < len(matches) else len(clean)
        chunks.append((start, match.group("date"), clean[start:end].strip()))
    return chunks


def extract_outage_type(text: str) -> str:
    upper = text.upper()
    if "TRIPPED ON FAULT" in upper:
        return "forced"
    if "LINE CLEAR AVAILED" in upper:
        return "planned_or_line_clear"
    if "LOAD CURTAILMENT" in upper:
        return "load_curtailment"
    if "SHUTDOWN" in upper:
        return "shutdown"
    return "other"


def extract_outage_remarks(text: str) -> str:
    upper = text.upper()
    for token in ["TRIPPED ON FAULT", "LINE CLEAR AVAILED", "LOAD CURTAILMENT", "SHUTDOWN"]:
        if token in upper:
            return token
    return ""


def parse_outage_events() -> pd.DataFrame:
    events: list[dict] = []

    for pdf_path in sorted(OUTAGE_PDF_DIR.glob("*.pdf")):
        reader = PdfReader(str(pdf_path))
        for page_num, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            for _, date_raw, chunk in parse_outage_chunks(text):
                head = ROW_START_RE.match(chunk)
                if not head:
                    continue

                serial_no = int(head.group("serial"))
                body = chunk[head.end() :].strip()
                zone = ""
                voltage_kv = None
                asset_text = body

                zone_match = ZONE_RE.match(body)
                if zone_match:
                    zone = zone_match.group("zone").strip()
                    voltage_kv = int(zone_match.group("voltage"))
                    asset_text = zone_match.group("rest").strip()

                times = TIME_RE.findall(asset_text)
                opened_at = normalize_time(times[0]) if len(times) >= 1 else None
                closed_at = normalize_time(times[1]) if len(times) >= 2 else None

                duration_minutes = None
                if len(times) >= 2:
                    second = normalize_time(times[1]) or ""
                    tail = asset_text.split(times[1], 1)[1] if times[1] in asset_text else ""
                    duration_match = re.search(r"\b(\d{1,2}:\d{2}(?::\d{2})?)\b", tail)
                    duration_minutes = parse_duration_minutes(duration_match.group(1)) if duration_match else None
                if duration_minutes is None:
                    duration_match = re.search(r"\b(\d{1,2}:\d{2}(?::\d{2})?)\b", chunk)
                    duration_minutes = parse_duration_minutes(duration_match.group(1)) if duration_match else None

                try:
                    outage_date = datetime.strptime(date_raw, "%d-%b-%y").date()
                except ValueError:
                    continue

                start_ts = None
                end_ts = None
                if opened_at:
                    start_ts = parse_time_to_datetime(outage_date, opened_at)
                if closed_at and start_ts is not None:
                    end_ts = parse_time_to_datetime(outage_date, closed_at)
                    if end_ts is not None and end_ts < start_ts:
                        end_ts += timedelta(days=1)
                elif start_ts is not None and duration_minutes is not None:
                    end_ts = start_ts + timedelta(minutes=float(duration_minutes))

                events.append(
                    {
                        "source_pdf": pdf_path.name,
                        "source_page": page_num,
                        "serial_no": serial_no,
                        "date": outage_date,
                        "transmission_zone": zone,
                        "voltage_class_kv": voltage_kv,
                        "asset_station_text": asset_text,
                        "opened_at": opened_at,
                        "closed_at": closed_at,
                        "duration_minutes": duration_minutes,
                        "outage_type": extract_outage_type(chunk),
                        "remarks": extract_outage_remarks(chunk),
                        "start_ts": start_ts,
                        "end_ts": end_ts,
                    }
                )

    df = pd.DataFrame(events)
    if df.empty:
        raise RuntimeError(f"No outage events parsed from: {OUTAGE_PDF_DIR}")
    df["is_high_voltage"] = df["voltage_class_kv"].fillna(0) >= 110
    return df


def build_hourly_outage_features(events: pd.DataFrame) -> pd.DataFrame:
    started = events.dropna(subset=["start_ts"]).copy()
    started["hour_ts"] = started["start_ts"].dt.floor("h")

    started_agg = (
        started.groupby("hour_ts")
        .agg(
            outages_started_hour=("serial_no", "count"),
            forced_outages_started_hour=("outage_type", lambda s: (s == "forced").sum()),
            line_clear_started_hour=("outage_type", lambda s: (s == "planned_or_line_clear").sum()),
            load_curtailment_started_hour=("outage_type", lambda s: (s == "load_curtailment").sum()),
            unique_zones_started_hour=("transmission_zone", pd.Series.nunique),
            high_voltage_started_hour=("is_high_voltage", "sum"),
            avg_duration_started_min_hour=("duration_minutes", "mean"),
        )
        .reset_index()
    )

    overlap_rows: list[dict] = []
    active_events = events.dropna(subset=["start_ts", "end_ts"]).copy()
    for event_id, row in active_events.reset_index(drop=True).iterrows():
        start_ts = row["start_ts"]
        end_ts = row["end_ts"]
        if end_ts <= start_ts:
            continue

        hour_cursor = start_ts.floor("h")
        while hour_cursor < end_ts:
            hour_end = hour_cursor + timedelta(hours=1)
            overlap_start = max(start_ts, hour_cursor)
            overlap_end = min(end_ts, hour_end)
            overlap_minutes = (overlap_end - overlap_start).total_seconds() / 60.0
            if overlap_minutes > 0:
                overlap_rows.append(
                    {
                        "hour_ts": hour_cursor,
                        "event_id": event_id,
                        "outage_type": row["outage_type"],
                        "zone": row["transmission_zone"],
                        "overlap_minutes": overlap_minutes,
                    }
                )
            hour_cursor = hour_end

    if overlap_rows:
        overlap = pd.DataFrame(overlap_rows)
        active_agg = (
            overlap.groupby("hour_ts")
            .agg(
                outages_active_hour=("event_id", pd.Series.nunique),
                outage_active_minutes_hour=("overlap_minutes", "sum"),
                forced_outage_active_minutes_hour=("overlap_minutes", lambda s: s[overlap.loc[s.index, "outage_type"] == "forced"].sum()),
                active_zones_hour=("zone", pd.Series.nunique),
            )
            .reset_index()
        )
    else:
        active_agg = pd.DataFrame(columns=["hour_ts", "outages_active_hour", "outage_active_minutes_hour", "forced_outage_active_minutes_hour", "active_zones_hour"])

    hourly = started_agg.merge(active_agg, on="hour_ts", how="outer").sort_values("hour_ts")
    return hourly


def build_daily_outage_features(events: pd.DataFrame) -> pd.DataFrame:
    base = (
        events.groupby("date")
        .agg(
            outages_total_day=("serial_no", "count"),
            forced_outages_day=("outage_type", lambda s: (s == "forced").sum()),
            line_clear_day=("outage_type", lambda s: (s == "planned_or_line_clear").sum()),
            load_curtailment_day=("outage_type", lambda s: (s == "load_curtailment").sum()),
            zones_affected_day=("transmission_zone", pd.Series.nunique),
            high_voltage_outages_day=("is_high_voltage", "sum"),
            avg_outage_duration_min_day=("duration_minutes", "mean"),
        )
        .reset_index()
    )

    zone_counts = (
        events.groupby(["date", "transmission_zone"])
        .size()
        .reset_index(name="zone_outage_count")
        .sort_values(["date", "zone_outage_count"], ascending=[True, False])
    )
    top_zone = zone_counts.groupby("date").head(1).rename(columns={"transmission_zone": "top_outage_zone_day", "zone_outage_count": "top_outage_zone_count_day"})

    zone_share = zone_counts.merge(base[["date", "outages_total_day"]], on="date", how="left")
    zone_share["share"] = zone_share["zone_outage_count"] / zone_share["outages_total_day"]
    concentration = zone_share.groupby("date")["share"].apply(lambda s: float((s**2).sum())).reset_index(name="zone_outage_concentration_hhi")

    daily = (
        base.merge(top_zone[["date", "top_outage_zone_day", "top_outage_zone_count_day"]], on="date", how="left")
        .merge(concentration, on="date", how="left")
    )
    return daily


def add_recommendation_features(master: pd.DataFrame) -> pd.DataFrame:
    # Use only dates with outage data to avoid false zeros in missing-source months.
    master["outage_source_month_available"] = master["month"].isin(set(master.loc[master["outages_total_day"].notna(), "month"]))

    metric_cols = [
        "outages_started_hour",
        "forced_outages_started_hour",
        "line_clear_started_hour",
        "load_curtailment_started_hour",
        "unique_zones_started_hour",
        "high_voltage_started_hour",
        "avg_duration_started_min_hour",
        "outages_active_hour",
        "outage_active_minutes_hour",
        "forced_outage_active_minutes_hour",
        "active_zones_hour",
        "outages_total_day",
        "forced_outages_day",
        "line_clear_day",
        "load_curtailment_day",
        "zones_affected_day",
        "high_voltage_outages_day",
        "avg_outage_duration_min_day",
        "top_outage_zone_count_day",
        "zone_outage_concentration_hhi",
    ]
    for col in metric_cols:
        if col in master.columns:
            master[col] = master[col].where(~master["outage_source_month_available"], master[col].fillna(0))

    master["is_evening_peak_window"] = master["hour"].between(18, 22)
    master["ev_shift_potential_mw"] = (master["load_2024_mw"] - master["daily_avg_load_2024_mw"]).clip(lower=0)

    daily_outage_p75 = master.groupby("date")["outages_active_hour"].transform(
        lambda s: np.nan if s.isna().all() else s.quantile(0.75)
    )
    ramp_scaled = master["ramp_1h_mw"].abs() / master["daily_peak_load_2024_mw"]
    outage_scaled = np.where(
        master["outage_source_month_available"],
        master["outages_active_hour"] / (daily_outage_p75.replace(0, np.nan)),
        0,
    )
    outage_scaled = np.clip(np.nan_to_num(outage_scaled, nan=0.0, posinf=1.0, neginf=0.0), 0, 1.5)
    master["grid_stress_index_0_100"] = (
        100
        * (
            0.60 * master["load_percentile_within_day"].fillna(0)
            + 0.25 * ramp_scaled.fillna(0).clip(0, 1)
            + 0.15 * outage_scaled
        )
    ).round(2)

    recommend = np.where(
        (master["load_percentile_within_day"] <= 0.35)
        & ((master["outages_active_hour"].fillna(0) <= daily_outage_p75.fillna(0)) | ~master["outage_source_month_available"]),
        "recommended",
        np.where(
            (master["load_percentile_within_day"] >= 0.8)
            | (master["outages_active_hour"].fillna(0) > daily_outage_p75.fillna(np.inf)),
            "avoid_peak",
            "neutral",
        ),
    )
    master["charging_window_recommendation"] = recommend
    return master


def style_output(path: Path) -> None:
    wb = load_workbook(path)
    ws = wb.active
    ws.title = "EV_Grid_EDA_2024"

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    widths = {
        "A": 12,
        "B": 19,
        "E": 11,
        "I": 11,
        "L": 13,
        "M": 14,
        "N": 13,
        "P": 16,
        "Q": 16,
        "R": 22,
        "S": 19,
        "T": 19,
        "U": 18,
        "V": 16,
        "W": 14,
        "X": 14,
        "Y": 16,
        "Z": 18,
        "AA": 16,
        "AB": 18,
        "AC": 18,
        "AD": 16,
        "AE": 16,
        "AF": 18,
        "AG": 18,
        "AH": 18,
        "AI": 17,
        "AJ": 18,
        "AK": 18,
        "AL": 24,
        "AM": 20,
        "AN": 24,
        "AO": 18,
        "AP": 20,
        "AQ": 18,
        "AR": 22,
        "AS": 19,
        "AT": 20,
    }
    for idx, _ in enumerate(ws[1], start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(letter, 12)

    for row in ws.iter_rows(min_row=2, max_col=ws.max_column):
        row[0].number_format = "yyyy-mm-dd"
        row[1].number_format = "yyyy-mm-dd hh:mm"
        for idx in [12, 13, 14, 15, 18, 19, 20, 24, 30, 39]:
            if idx < len(row):
                row[idx].number_format = "0.00"
        for idx in [16, 21]:
            if idx < len(row):
                row[idx].number_format = "0.00%"

    wb.save(path)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    load_df = build_load_hourly_dataset()
    outage_events = parse_outage_events()
    outage_hourly = build_hourly_outage_features(outage_events)
    outage_daily = build_daily_outage_features(outage_events)

    master = (
        load_df.merge(outage_hourly, left_on="timestamp", right_on="hour_ts", how="left")
        .drop(columns=["hour_ts"], errors="ignore")
        .merge(outage_daily, on="date", how="left")
    )
    master = add_recommendation_features(master)

    ordered_cols = [
        "date",
        "timestamp",
        "year",
        "month",
        "month_name",
        "quarter",
        "day",
        "day_of_year",
        "day_of_week",
        "is_weekend",
        "hour",
        "time_block",
        "load_2024_mw",
        "load_2023_mw",
        "frequency_hz",
        "load_delta_2024_vs_2023_mw",
        "load_delta_2024_vs_2023_pct",
        "daily_peak_load_2024_mw",
        "daily_min_load_2024_mw",
        "daily_avg_load_2024_mw",
        "load_factor_vs_daily_peak",
        "is_daily_peak_hour",
        "is_daily_min_hour",
        "ramp_1h_mw",
        "load_percentile_within_day",
        "is_evening_peak_window",
        "ev_shift_potential_mw",
        "outage_source_month_available",
        "outages_started_hour",
        "outages_active_hour",
        "outage_active_minutes_hour",
        "forced_outages_started_hour",
        "forced_outage_active_minutes_hour",
        "line_clear_started_hour",
        "load_curtailment_started_hour",
        "unique_zones_started_hour",
        "active_zones_hour",
        "high_voltage_started_hour",
        "avg_duration_started_min_hour",
        "outages_total_day",
        "forced_outages_day",
        "line_clear_day",
        "load_curtailment_day",
        "zones_affected_day",
        "high_voltage_outages_day",
        "avg_outage_duration_min_day",
        "top_outage_zone_day",
        "top_outage_zone_count_day",
        "zone_outage_concentration_hhi",
        "grid_stress_index_0_100",
        "charging_window_recommendation",
        "load_curve_source_file",
        "load_data_coverage_note",
    ]

    master = master[ordered_cols]
    master.to_excel(OUTPUT_XLSX, index=False, engine="openpyxl")
    master.to_csv(OUTPUT_CSV, index=False)
    style_output(OUTPUT_XLSX)

    missing_load_note = master.loc[master["load_data_coverage_note"] != "", "load_data_coverage_note"].head(1).tolist()
    month_coverage = sorted(master.loc[master["outage_source_month_available"], "month"].unique().tolist())

    print(f"Output: {OUTPUT_XLSX}")
    print(f"CSV: {OUTPUT_CSV}")
    print(f"Rows: {len(master)}")
    print(f"Hours per available day (min/max): {master.groupby('date')['hour'].count().min()}/{master.groupby('date')['hour'].count().max()}")
    print(f"Dates in load data: {master['date'].nunique()}")
    print(f"Outage months available: {month_coverage}")
    print(f"Charging recommendation counts: {master['charging_window_recommendation'].value_counts().to_dict()}")
    if missing_load_note:
        print(f"Load note: {missing_load_note[0]}")


if __name__ == "__main__":
    main()
